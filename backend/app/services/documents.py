import csv
import io
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_document(filename: str, content: bytes) -> tuple[str, int]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8", errors="replace"), 1
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(content.decode("utf-8", errors="replace"))))
        return "\n".join(" | ".join(row) for row in rows), 1
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)
    if suffix == ".docx":
        document = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs), 1
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"# Sheet: {sheet.title}")
            lines.extend(" | ".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True))
        return "\n".join(lines), len(workbook.worksheets)
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")


def chunk_text(text: str, size: int, overlap: int) -> list[dict]:
    if overlap >= size:
        raise ValueError("overlap must be smaller than size")
    words = text.split()
    chunks, start, index = [], 0, 0
    while start < len(words):
        end = min(start + size, len(words))
        chunks.append({"id": f"chunk-{index:04d}", "text": " ".join(words[start:end]), "start": start, "end": end, "metadata": {"tokens_estimate": end - start}})
        if end == len(words):
            break
        start = end - overlap
        index += 1
    return chunks
